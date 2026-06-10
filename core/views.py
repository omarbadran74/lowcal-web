from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponseForbidden, HttpResponse
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import ensure_csrf_cookie
from django.contrib import messages
from django.utils import timezone
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .models import (
    User, FoodCategory, FoodItem, SubscriptionPlan,
    Subscriber, DailyMealSelection, DailySubmission, DailyDelivery, MEAL_TYPES
)


# -- Auth --

def login_view(request):
    if request.user.is_authenticated:
        return _redirect_by_role(request.user)
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user:
            login(request, user)
            return _redirect_by_role(user)
        messages.error(request, 'Invalid username or password')
    return render(request, 'login.html')


def logout_view(request):
    logout(request)
    return redirect('login')


def _redirect_by_role(user):
    if user.is_admin_role():
        return redirect('admin_panel:dashboard')
    return redirect('cashier:dashboard')


# -- Admin Panel --

def admin_required(view_func):
    @login_required
    def wrapper(request, *args, **kwargs):
        if not request.user.is_admin_role():
            return HttpResponseForbidden('Access denied')
        return view_func(request, *args, **kwargs)
    wrapper.__name__ = view_func.__name__
    return wrapper


@admin_required
def admin_dashboard(request):
    context = {
        'total_subscribers': Subscriber.objects.count(),
        'total_food_items': FoodItem.objects.filter(is_active=True).count(),
        'total_categories': FoodCategory.objects.filter(is_active=True).count(),
        'recent_subscribers': Subscriber.objects.select_related('plan').order_by('-created_at')[:5],
    }
    return render(request, 'admin_panel/dashboard.html', context)


@admin_required
def category_list(request):
    categories = FoodCategory.objects.prefetch_related('items').order_by('meal_type', 'name')
    return render(request, 'admin_panel/category_list.html', {'categories': categories})


@admin_required
def category_create(request):
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        meal_type = request.POST.get('meal_type')
        if name and meal_type:
            FoodCategory.objects.create(name=name, meal_type=meal_type)
            messages.success(request, 'Category added successfully')
            return redirect('admin_panel:category_list')
        messages.error(request, 'Please fill all fields')
    return render(request, 'admin_panel/category_form.html', {
        'meal_types': MEAL_TYPES,
        'title': 'Add Category',
    })


@admin_required
def category_edit(request, pk):
    category = get_object_or_404(FoodCategory, pk=pk)
    if request.method == 'POST':
        category.name = request.POST.get('name', '').strip()
        category.meal_type = request.POST.get('meal_type')
        category.is_active = request.POST.get('is_active') == 'on'
        category.save()
        messages.success(request, 'Category updated successfully')
        return redirect('admin_panel:category_list')
    return render(request, 'admin_panel/category_form.html', {
        'category': category,
        'meal_types': MEAL_TYPES,
        'title': 'Edit Category',
    })


@admin_required
@require_POST
def category_delete(request, pk):
    category = get_object_or_404(FoodCategory, pk=pk)
    category.is_active = False
    category.save()
    messages.success(request, 'Category deleted')
    return redirect('admin_panel:category_list')


@admin_required
def item_list(request):
    category_id = request.GET.get('category')
    items = FoodItem.objects.prefetch_related('categories').filter(is_active=True)
    if category_id:
        items = items.filter(categories__id=category_id)
    categories = FoodCategory.objects.filter(is_active=True)
    return render(request, 'admin_panel/item_list.html', {
        'items': items,
        'categories': categories,
        'selected_category': int(category_id) if category_id else None,
    })


@admin_required
def item_create(request):
    categories = FoodCategory.objects.filter(is_active=True).order_by('meal_type', 'name')
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        category_ids = request.POST.getlist('categories')
        description = request.POST.get('description', '').strip()
        calories = request.POST.get('calories') or None
        if name and category_ids:
            item = FoodItem.objects.create(
                name=name,
                description=description,
                calories=int(calories) if calories else None,
            )
            item.categories.set(category_ids)
            messages.success(request, 'Dish added successfully')
            return redirect('admin_panel:item_list')
        messages.error(request, 'Please fill all required fields')
    return render(request, 'admin_panel/item_form.html', {
        'categories': categories,
        'selected_categories': [],
        'title': 'Add Dish',
    })


@admin_required
def item_edit(request, pk):
    item = get_object_or_404(FoodItem, pk=pk)
    categories = FoodCategory.objects.filter(is_active=True).order_by('meal_type', 'name')
    if request.method == 'POST':
        item.name = request.POST.get('name', '').strip()
        item.description = request.POST.get('description', '').strip()
        calories = request.POST.get('calories') or None
        item.calories = int(calories) if calories else None
        item.is_active = request.POST.get('is_active') == 'on'
        item.save()
        item.categories.set(request.POST.getlist('categories'))
        messages.success(request, 'Dish updated successfully')
        return redirect('admin_panel:item_list')
    return render(request, 'admin_panel/item_form.html', {
        'item': item,
        'categories': categories,
        'selected_categories': list(item.categories.values_list('id', flat=True)),
        'title': 'Edit Dish',
    })


@admin_required
@require_POST
def item_delete(request, pk):
    item = get_object_or_404(FoodItem, pk=pk)
    item.is_active = False
    item.save()
    messages.success(request, 'Dish deleted')
    return redirect('admin_panel:item_list')


@admin_required
def plan_list(request):
    plans = SubscriptionPlan.objects.all().order_by('goal', 'meals_per_day', 'days')
    return render(request, 'admin_panel/plan_list.html', {'plans': plans})


@admin_required
def plan_create(request):
    if request.method == 'POST':
        goal = request.POST.get('goal')
        meals_per_day = request.POST.get('meals_per_day')
        days = request.POST.get('days')
        price = request.POST.get('price')
        if all([goal, meals_per_day, days, price]):
            SubscriptionPlan.objects.get_or_create(
                goal=goal, meals_per_day=int(meals_per_day), days=int(days),
                defaults={'price': price}
            )
            messages.success(request, 'Plan added successfully')
            return redirect('admin_panel:plan_list')
    return render(request, 'admin_panel/plan_form.html', {'title': 'Add Plan'})


@admin_required
def plan_edit(request, pk):
    plan = get_object_or_404(SubscriptionPlan, pk=pk)
    if request.method == 'POST':
        plan.goal = request.POST.get('goal')
        plan.meals_per_day = int(request.POST.get('meals_per_day'))
        plan.days = int(request.POST.get('days'))
        plan.price = request.POST.get('price')
        plan.save()
        messages.success(request, 'Plan updated successfully')
        return redirect('admin_panel:plan_list')
    return render(request, 'admin_panel/plan_form.html', {'plan': plan, 'title': 'Edit Plan'})


@admin_required
@require_POST
def plan_delete(request, pk):
    plan = get_object_or_404(SubscriptionPlan, pk=pk)
    plan.delete()
    messages.success(request, 'Plan deleted')
    return redirect('admin_panel:plan_list')


@admin_required
def admin_subscribers(request):
    subscribers = Subscriber.objects.select_related('plan', 'created_by').order_by('-created_at')
    return render(request, 'admin_panel/subscribers.html', {'subscribers': subscribers})


@admin_required
@require_POST
def reset_day(request, pk, day_number):
    subscriber = get_object_or_404(Subscriber, pk=pk)
    DailySubmission.objects.filter(subscriber=subscriber, day_number=day_number).delete()
    DailyDelivery.objects.filter(subscriber=subscriber, day_number=day_number).delete()
    messages.success(request, f'Day {day_number} reset — subscriber can now reselect and resubmit')
    return redirect('cashier:subscriber_detail', pk=pk)


@admin_required
@require_POST
def add_absent_day(request, pk):
    subscriber = get_object_or_404(Subscriber, pk=pk)
    subscriber.custom_days = subscriber.total_days() + 1
    subscriber.save(update_fields=['custom_days'])
    messages.success(request, f'Absent day added — total is now {subscriber.total_days()} days')
    return redirect('cashier:subscriber_detail', pk=pk)


@admin_required
@require_POST
def subscriber_delete(request, pk):
    subscriber = get_object_or_404(Subscriber, pk=pk)
    name = subscriber.name
    subscriber.delete()
    messages.success(request, f'Subscriber {name} deleted')
    return redirect('cashier:dashboard')


# -- Cashier Panel --

def cashier_required(view_func):
    @login_required
    def wrapper(request, *args, **kwargs):
        if not (request.user.is_cashier_role() or request.user.is_admin_role()):
            return HttpResponseForbidden('Access denied')
        return view_func(request, *args, **kwargs)
    wrapper.__name__ = view_func.__name__
    return wrapper


@cashier_required
def cashier_dashboard(request):
    subscribers = Subscriber.objects.select_related('plan', 'created_by').prefetch_related('submissions').order_by('-created_at')
    return render(request, 'cashier/dashboard.html', {'subscribers': subscribers})


@cashier_required
def subscriber_create(request):
    plans = SubscriptionPlan.objects.all().order_by('goal', 'meals_per_day', 'days')
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        phone = request.POST.get('phone', '').strip()
        plan_id = request.POST.get('plan')
        meal_types = request.POST.getlist('meal_types')
        start_date = request.POST.get('start_date')
        notes = request.POST.get('notes', '').strip()
        raw_days = request.POST.get('custom_days', '').strip()
        raw_price = request.POST.get('custom_price', '').strip()
        raw_paid = request.POST.get('paid_amount', '').strip()
        paid_by = request.POST.get('paid_by', '').strip()
        if name and phone and plan_id and meal_types and start_date:
            plan = SubscriptionPlan.objects.get(pk=plan_id)
            cd = int(raw_days) if raw_days and raw_days.isdigit() else None
            custom_days = cd if cd and cd != plan.days else None
            try:
                cp = float(raw_price) if raw_price else None
                custom_price = cp if cp and cp != float(plan.price) else None
            except ValueError:
                custom_price = None
            try:
                paid_amount = float(raw_paid) if raw_paid else None
            except ValueError:
                paid_amount = None
            subscriber = Subscriber.objects.create(
                name=name, phone=phone, plan=plan,
                meal_types=','.join(meal_types),
                start_date=start_date, notes=notes, created_by=request.user,
                custom_days=custom_days, custom_price=custom_price,
                paid_amount=paid_amount, paid_by=paid_by,
            )
            messages.success(request, f'Subscriber {name} created successfully!')
            return redirect('cashier:subscriber_detail', pk=subscriber.pk)
        messages.error(request, 'Please fill all required fields')
    return render(request, 'cashier/subscriber_form.html', {
        'plans': plans,
        'meal_types': MEAL_TYPES,
        'title': 'New Subscriber',
    })


@admin_required
def subscriber_edit(request, pk):
    subscriber = get_object_or_404(Subscriber, pk=pk)
    plans = SubscriptionPlan.objects.all().order_by('goal', 'meals_per_day', 'days')
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        phone = request.POST.get('phone', '').strip()
        plan_id = request.POST.get('plan')
        meal_types = request.POST.getlist('meal_types')
        start_date = request.POST.get('start_date')
        notes = request.POST.get('notes', '').strip()
        raw_days = request.POST.get('custom_days', '').strip()
        raw_price = request.POST.get('custom_price', '').strip()
        raw_paid = request.POST.get('paid_amount', '').strip()
        paid_by = request.POST.get('paid_by', '').strip()
        if name and phone and plan_id and meal_types and start_date:
            plan = SubscriptionPlan.objects.get(pk=plan_id)
            cd = int(raw_days) if raw_days and raw_days.isdigit() else None
            custom_days = cd if cd and cd != plan.days else None
            try:
                cp = float(raw_price) if raw_price else None
                custom_price = cp if cp and cp != float(plan.price) else None
            except ValueError:
                custom_price = None
            try:
                paid_amount = float(raw_paid) if raw_paid else None
            except ValueError:
                paid_amount = None
            subscriber.name = name
            subscriber.phone = phone
            subscriber.plan = plan
            subscriber.meal_types = ','.join(meal_types)
            subscriber.start_date = start_date
            subscriber.notes = notes
            subscriber.custom_days = custom_days
            subscriber.custom_price = custom_price
            subscriber.paid_amount = paid_amount
            subscriber.paid_by = paid_by
            subscriber.save()
            messages.success(request, f'Subscriber {name} updated successfully!')
            return redirect('cashier:subscriber_detail', pk=subscriber.pk)
        messages.error(request, 'Please fill all required fields')
    return render(request, 'cashier/subscriber_form.html', {
        'plans': plans,
        'meal_types': MEAL_TYPES,
        'subscriber': subscriber,
        'title': f'Edit — {subscriber.name}',
    })


@cashier_required
def subscriber_detail(request, pk):
    subscriber = get_object_or_404(Subscriber, pk=pk)
    submitted_days = set(
        DailySubmission.objects.filter(subscriber=subscriber).values_list('day_number', flat=True)
    )
    days_data = []
    for day in range(1, subscriber.total_days() + 1):
        selections = DailyMealSelection.objects.filter(
            subscriber=subscriber, day_number=day
        ).select_related('food_item')
        day_selections = {s.meal_type: s.food_item for s in selections}
        days_data.append({
            'day': day,
            'submitted': day in submitted_days,
            'selections': day_selections,
        })
    base_url = request.build_absolute_uri(f'/s/{subscriber.token}/')
    return render(request, 'cashier/subscriber_detail.html', {
        'subscriber': subscriber,
        'days_data': days_data,
        'subscriber_url': base_url,
        'meal_types': subscriber.get_meal_types_list(),
        'meal_types_display': dict(MEAL_TYPES),
    })


def _build_orders(day_filter=None, name_filter='', plan_filter=''):
    subs = Subscriber.objects.select_related('plan').prefetch_related('selections__food_item', 'submissions')
    if name_filter:
        subs = subs.filter(name__icontains=name_filter)
    if plan_filter:
        subs = subs.filter(plan__goal=plan_filter)

    delivered_set = set(DailyDelivery.objects.values_list('subscriber_id', 'day_number'))

    orders = []
    for sub in subs:
        submitted_days = set(sub.submissions.values_list('day_number', flat=True))
        for d in range(1, sub.total_days() + 1):
            if day_filter and str(d) != str(day_filter):
                continue
            if d not in submitted_days:
                continue
            sels = {s.meal_type: s.food_item.name if s.food_item else '-' for s in sub.selections.filter(day_number=d)}
            orders.append({
                'subscriber': sub,
                'day': d,
                'selections': sels,
                'meal_types': sub.get_meal_types_list(),
                'delivered': (sub.pk, d) in delivered_set,
            })
    return orders


@cashier_required
def orders_view(request):
    day_filter = request.GET.get('day', '')
    name_filter = request.GET.get('name', '').strip()
    plan_filter = request.GET.get('plan', '')

    orders = _build_orders(day_filter or None, name_filter, plan_filter)

    return render(request, 'cashier/orders.html', {
        'orders': orders,
        'meal_types_display': dict(MEAL_TYPES),
        'selected_day': day_filter,
        'name_filter': name_filter,
        'plan_filter': plan_filter,
        'goal_choices': SubscriptionPlan.GOAL_CHOICES,
    })


@cashier_required
@require_POST
def toggle_delivery(request, subscriber_id, day_number):
    subscriber = get_object_or_404(Subscriber, pk=subscriber_id)
    delivery, created = DailyDelivery.objects.get_or_create(
        subscriber=subscriber,
        day_number=day_number,
        defaults={'delivered_by': request.user}
    )
    if not created:
        delivery.delete()
        return JsonResponse({'delivered': False})
    return JsonResponse({'delivered': True})


@cashier_required
def orders_export_excel(request):
    day_filter = request.GET.get('day')
    orders = _build_orders(day_filter)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Orders'

    header_fill = PatternFill(start_color='FF6B00', end_color='FF6B00', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=12)
    center = Alignment(horizontal='center', vertical='center')

    headers = ['Name', 'Phone', 'Plan', 'Day', 'Breakfast', 'Lunch', 'Dinner', 'Snack']
    meal_type_cols = {'breakfast': 4, 'lunch': 5, 'dinner': 6, 'snack': 7}

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        ws.column_dimensions[cell.column_letter].width = 20

    for row_idx, order in enumerate(orders, 2):
        sub = order['subscriber']
        ws.cell(row=row_idx, column=1, value=sub.name)
        ws.cell(row=row_idx, column=2, value=sub.phone)
        ws.cell(row=row_idx, column=3, value=str(sub.plan))
        ws.cell(row=row_idx, column=4, value=order['day'])
        for mt, col in meal_type_cols.items():
            if mt in order['meal_types']:
                ws.cell(row=row_idx, column=col + 1, value=order['selections'].get(mt, ''))

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    day_str = f'_day{day_filter}' if day_filter else ''
    resp['Content-Disposition'] = f'attachment; filename="orders{day_str}.xlsx"'
    wb.save(resp)
    return resp


# -- Subscriber Public Page --

@ensure_csrf_cookie
def subscriber_page(request, token):
    subscriber = get_object_or_404(Subscriber, token=token)
    plan = subscriber.plan
    meal_types_list = subscriber.get_meal_types_list()

    submitted_days = set(
        DailySubmission.objects.filter(subscriber=subscriber).values_list('day_number', flat=True)
    )

    existing_selections = {}
    for sel in DailyMealSelection.objects.filter(subscriber=subscriber).select_related('food_item'):
        existing_selections[(sel.day_number, sel.meal_type)] = sel

    food_items_by_type = {}
    for mt in meal_types_list:
        food_items_by_type[mt] = list(
            FoodItem.objects.filter(categories__meal_type=mt, is_active=True)
            .distinct()
            .values('id', 'name', 'calories', 'description')
        )

    today = timezone.localdate()
    days_elapsed = 0
    if subscriber.start_date:
        days_elapsed = max(0, (today - subscriber.start_date).days)

    days_data = []
    for day in range(1, subscriber.total_days() + 1):
        submitted = day in submitted_days
        is_past = (day <= days_elapsed) and not submitted
        meals = []
        all_selected = True
        for mt in meal_types_list:
            sel = existing_selections.get((day, mt))
            selected_item = sel.food_item if sel else None
            if not selected_item:
                all_selected = False
            meals.append({
                'meal_type': mt,
                'meal_type_display': dict(MEAL_TYPES).get(mt, mt),
                'selected_item': selected_item,
                'food_items': food_items_by_type.get(mt, []),
            })
        days_data.append({
            'day': day,
            'submitted': submitted,
            'is_past': is_past,
            'meals': meals,
            'can_submit': all_selected and not submitted and not is_past,
        })

    return render(request, 'subscriber/page.html', {
        'subscriber': subscriber,
        'plan': plan,
        'days_data': days_data,
        'meal_types_display': dict(MEAL_TYPES),
        'submitted_count': len(submitted_days),
        'days_elapsed': days_elapsed,
        'current_day': min(days_elapsed + 1, subscriber.total_days()),
        'remaining_days': max(0, subscriber.total_days() - days_elapsed),
    })


@require_POST
def subscriber_select_meal(request, token):
    subscriber = get_object_or_404(Subscriber, token=token)
    data = json.loads(request.body)
    day_number = int(data.get('day'))
    meal_type = data.get('meal_type')
    food_item_id = data.get('food_item_id')

    if DailySubmission.objects.filter(subscriber=subscriber, day_number=day_number).exists():
        return JsonResponse({'error': 'This day was already submitted'}, status=400)

    if meal_type not in subscriber.get_meal_types_list():
        return JsonResponse({'error': 'Invalid meal type'}, status=400)

    food_item = get_object_or_404(FoodItem, pk=food_item_id, is_active=True)

    DailyMealSelection.objects.update_or_create(
        subscriber=subscriber,
        day_number=day_number,
        meal_type=meal_type,
        defaults={'food_item': food_item},
    )
    return JsonResponse({'success': True, 'item_name': food_item.name})


@require_POST
def subscriber_submit_day(request, token):
    subscriber = get_object_or_404(Subscriber, token=token)
    data = json.loads(request.body)
    day_number = int(data.get('day'))

    if DailySubmission.objects.filter(subscriber=subscriber, day_number=day_number).exists():
        return JsonResponse({'error': 'This day was already submitted'}, status=400)

    meal_types_list = subscriber.get_meal_types_list()
    selected_count = DailyMealSelection.objects.filter(
        subscriber=subscriber,
        day_number=day_number,
        meal_type__in=meal_types_list,
        food_item__isnull=False,
    ).count()

    if selected_count < len(meal_types_list):
        return JsonResponse({'error': 'Please select all meals first'}, status=400)

    DailySubmission.objects.create(subscriber=subscriber, day_number=day_number)
    return JsonResponse({'success': True})
